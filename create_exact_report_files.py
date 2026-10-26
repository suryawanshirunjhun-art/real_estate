"""
Generate the exact formal Word document (.docx) adhering to the university template:
- Margins: 1 inch
- Times New Roman font throughout
- Project Title: 16 pt Bold
- Headings / Subheadings: 14 pt Bold
- Normal text: 12 pt Justified
- Table formatting and image embeds
"""
import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCX_FILENAME = "SI_Lab_Project_Report_Runjhun_Suryawanshi.docx"

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_docx():
    doc = Document()

    # Set 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Set default style to Times New Roman 12pt
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # Helper function for adding paragraphs with formatting
    def add_p(text, font_size=12, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(font_size)
        run.bold = bold
        run.italic = italic
        return p

    def add_heading1(text):
        return add_p(text, font_size=14, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=12, space_after=4)

    def add_heading2(text):
        return add_p(text, font_size=14, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=8, space_after=3)

    # =========================================================================
    # PAGE 1: COVER PAGE
    # =========================================================================
    add_p("RAMDEOBABA UNIVERSITY, NAGPUR", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", font_size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_p("B.Tech. CSE (Artificial Intelligence and Data Science)\n(AIDS)", font_size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)

    # Logo
    if os.path.exists("university_logo.jpeg"):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_after = Pt(20)
        p_logo.add_run().add_picture("university_logo.jpeg", width=Inches(1.8))

    add_p("SEMESTER V", font_size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_p("SI LAB", font_size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_p("PROJECT REPORT", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)

    # Project Title
    add_p("REAL ESTATE PRICE PREDICTION AND MARKET VALUATION ANALYSIS USING MACHINE LEARNING", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)

    add_p("Submitted by", font_size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    # Student Table
    table_student = doc.add_table(rows=2, cols=3)
    table_student.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_widths = [Inches(2.5), Inches(1.2), Inches(2.3)]

    headers = ["Student Name", "Roll No.", "Enrollment No."]
    for i, h in enumerate(headers):
        cell = table_student.cell(0, i)
        cell.width = col_widths[i]
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, 120, 120, 150, 150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)

    row1_data = ["Runjhun Suryawanshi", "18", "[Enrollment No.]"]
    for i, d in enumerate(row1_data):
        cell = table_student.cell(1, i)
        cell.width = col_widths[i]
        set_cell_margins(cell, 120, 120, 150, 150)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(d)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)

    # Table borders
    tblPr = table_student._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="0F172A"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="0F172A"/>'
        f'<w:left w:val="single" w:sz="6" w:space="0" w:color="0F172A"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="0F172A"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(30)
    add_p("Academic Year: 2026–27", font_size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)

    # Page Break to Page 2
    doc.add_page_break()

    # =========================================================================
    # PAGE 2: SECTIONS 1, 2, 3
    # =========================================================================
    add_p("REAL ESTATE PRICE PREDICTION AND MARKET VALUATION ANALYSIS", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    add_p("SI Lab Project Report", font_size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    add_heading1("1. PROJECT OBJECTIVE")
    add_p(
        "The primary purpose of this project is to develop an automated, end-to-end Machine Learning regression and market analysis system (EstatePulse AI) to forecast residential real estate prices based on structural, physical, and geographic attributes. The system eliminates subjective pricing biases, providing institutional investors, appraisers, and homebuyers with sub-second property appraisals alongside statistical confidence bounds."
    )
    add_p("Specific objectives for the project:", bold=True)
    add_p("● Load and explore the multi-metro residential housing transaction dataset.")
    add_p("● Preprocess the data through median/mode imputation, duplicate auditing, One-Hot encoding, and feature scaling.")
    add_p("● Extract and evaluate relevant structural and spatial features using correlation matrix analysis.")
    add_p("● Train multiple regression models: Multiple Linear Regression, Ridge (L2), Lasso (L1), Decision Tree, and Random Forest.")
    add_p("● Evaluate and compare model performance using R², Adjusted R², MAE, RMSE, MAPE, and 5-fold cross-validation.")
    add_p("● Select the best model based strictly on cross-validated performance to prevent data leakage.")
    add_p("● Develop an interactive multi-page GUI using Streamlit to accept input attributes and display live valuations and scenario simulations.")

    add_heading1("2. INTRODUCTION")
    add_p(
        "Real estate valuation has traditionally relied on subjective manual appraisals, broker surveys, or rigid linear metrics (e.g., fixed rate per square foot). These conventional methods fail because property pricing involves complex non-linear interactions between physical features (carpet area, bedrooms), structural condition (building age, interior fitout), and spatial factors (city economic baselines, transit distance to central business districts). Machine Learning offers a rigorous data-driven solution by learning these multi-dimensional interactions simultaneously."
    )

    add_heading2("2.1 Problem Statement")
    add_p(
        "To accurately estimate the fair market valuation of residential properties given multi-dimensional physical and geographic inputs, minimizing forecasting error (RMSE and MAPE) and preventing overfitting through rigorous cross-validation."
    )

    add_heading2("2.2 Scope of the Project")
    add_p(
        "Functions included in the project comprise automated dataset ingestion across 5 Indian metro corridors (Mumbai, Bengaluru, Pune, Nagpur, Raipur), statistical profiling, One-Hot encoding, 5-fold cross-validation grid search, residual diagnostics, and an interactive multi-page web GUI with preset archetypes and live CapEx sensitivity simulation. Limitations: Valuations assume standard market conditions and cannot account for extreme localized macroeconomic shocks or subjective architectural heritage prestige not captured in numerical attributes."
    )

    add_heading1("3. PROJECT METHODOLOGY")
    add_p(
        "The project follows a standard machine learning lifecycle: Data Ingestion → Preprocessing & Imputation → Exploratory Data Analysis & Correlation Auditing → 80/20 Train-Test Split → 5-Fold Cross-Validation Grid Search → Algorithmic Benchmarking → Model Serialization → Interactive GUI Deployment."
    )

    add_heading2("3.1 Dataset Loading and Exploration")
    add_p("● Dataset Name & Source: Indian Metro Residential Real Estate Benchmark Dataset (compiled from housing transactions across major Indian urban centers).")
    add_p("● Dimensions: 1,515 observations, 9 features (8 independent predictors + 1 continuous target variable).")
    add_p("● Target Variable: price (Continuous numerical in Indian Rupees, ranging from ₹15 Lakhs to ₹7.5+ Crores).")
    add_p("● Features: area_sqft (300–4,000 sq.ft), bedrooms (1–6 BHK), bathrooms (1–6), age_years (0–30 yrs), distance_to_center_km (0.5–25 km), city (Mumbai, Pune, Bengaluru, Nagpur, Raipur), property_type (Apartment, House, Villa), furnishing (Unfurnished, Semi-Furnished, Furnished).")

    add_heading2("3.2 Data Preprocessing")
    add_p("● Handling Missing Values: Numeric features imputed using median imputation via SimpleImputer(strategy='median'); categorical features imputed using mode imputation (strategy='most_frequent').")
    add_p("● Duplicate Auditing: Duplicate records audited via df.duplicated().sum() and dropped to eliminate redundant sample bias.")
    add_p("● Categorical Encoding: Categorical features converted using OneHotEncoder(drop='first', sparse_output=False, handle_unknown='ignore').")
    add_p("● Feature Scaling: StandardScaler incorporated into pipeline for regularized linear algorithms (Ridge and Lasso).")
    add_p("● Train-Test Split: Partitioned into 80% Training Set (1,200 samples) and 20% Held-Out Testing Set (300 samples) using random_state=42.")

    add_heading2("3.3 Feature Extraction / Selection")
    add_p(
        "Correlation analysis confirmed that living area (r = +0.51) and Mumbai location index (r = +0.68) are the strongest linear drivers of price, while distance to CBD reflects spatial decay (r = -0.04). Lasso (L1) regularization verified that all 8 features contribute explanatory variance, and thus all available features were retained."
    )

    # Page Break to Page 3
    doc.add_page_break()

    # =========================================================================
    # PAGE 3: SECTIONS 3.4, 3.5, 3.6, 4.1, 4.2
    # =========================================================================
    add_heading1("3.4 Training Multiple Regression Models")
    add_p("Five supervised regression models were implemented and tuned using 5-fold cross-validation:")

    # Table 3.4
    table_models = doc.add_table(rows=6, cols=3)
    table_models.alignment = WD_TABLE_ALIGNMENT.CENTER
    m_widths = [Inches(0.8), Inches(2.2), Inches(3.5)]
    m_headers = ["Sr. No.", "Model Name", "Purpose / Notes"]
    for i, h in enumerate(m_headers):
        cell = table_models.cell(0, i)
        cell.width = m_widths[i]
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(10)

    m_rows = [
        ("1", "Multiple Linear Regression", "Baseline OLS model assuming strictly additive linear relationships across all predictors."),
        ("2", "Ridge Regression (L2)", "Penalizes sum of squared weights (alpha * sum(w_j^2)) to stabilize weights and mitigate multicollinearity."),
        ("3", "Lasso Regression (L1)", "Penalizes sum of absolute weights (alpha * sum(|w_j|)) to shrink uninformative coefficients strictly to zero."),
        ("4", "Decision Tree Regressor", "Non-parametric model recursively partitioning feature space into piecewise constant segments to model non-linearities."),
        ("5", "Random Forest Regressor", "Ensemble bagging model training 100 decorrelated decision trees on bootstrap subsets, averaging predictions to minimize variance.")
    ]
    for r_idx, row in enumerate(m_rows, start=1):
        for c_idx, val in enumerate(row):
            cell = table_models.cell(r_idx, c_idx)
            cell.width = m_widths[c_idx]
            set_cell_margins(cell, 60, 60, 100, 100)
            if r_idx % 2 == 0:
                set_cell_background(cell, "F8FAFC")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(val)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            if c_idx == 1:
                run.bold = True

    add_heading1("3.5 Performance Evaluation and Model Comparison")
    add_p("Models were evaluated across standardized regression metrics:")

    # Table 3.5: Benchmark
    table_bench = doc.add_table(rows=6, cols=5)
    table_bench.alignment = WD_TABLE_ALIGNMENT.CENTER
    b_widths = [Inches(2.0), Inches(1.1), Inches(1.1), Inches(1.2), Inches(1.1)]
    b_headers = ["Model", "Accuracy / R²", "MAE (₹)", "RMSE (₹)", "MAPE (%)"]
    for i, h in enumerate(b_headers):
        cell = table_bench.cell(0, i)
        cell.width = b_widths[i]
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(10)

    b_rows = [
        ("Random Forest", "96.04% (0.9604)", "₹ 7,83,302", "₹ 15,84,460", "8.95%"),
        ("Decision Tree", "94.65% (0.9465)", "₹ 10,13,071", "₹ 18,41,921", "11.35%"),
        ("Lasso Regression", "88.69% (0.8869)", "₹ 17,79,466", "₹ 26,78,266", "29.05%"),
        ("Ridge Regression", "88.69% (0.8869)", "₹ 17,79,462", "₹ 26,78,268", "29.05%"),
        ("Linear Regression", "88.69% (0.8869)", "₹ 17,79,475", "₹ 26,78,267", "29.06%")
    ]
    for r_idx, row in enumerate(b_rows, start=1):
        for c_idx, val in enumerate(row):
            cell = table_bench.cell(r_idx, c_idx)
            cell.width = b_widths[c_idx]
            set_cell_margins(cell, 60, 60, 100, 100)
            if r_idx == 1:
                set_cell_background(cell, "F0FDF4")
            elif r_idx % 2 == 0:
                set_cell_background(cell, "F8FAFC")
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            if r_idx == 1 or c_idx == 0:
                run.bold = True

    add_heading1("3.6 Best Model Selection")
    add_p(
        "Random Forest Regressor was selected as the champion model. It achieved an out-of-sample Test R² of 0.9604 (96.04% variance explained), reduced Mean Absolute Error by over 55% compared to linear baselines (₹7.83 Lakhs vs ₹17.79 Lakhs), and achieved a commercial-grade error rate of 8.95% MAPE (under the 10% appraisal standard). Furthermore, it captured multiplicative interactions (e.g., location multiplier * typology * area) that linear models cannot resolve."
    )

    add_heading1("4. PROJECT IMPLEMENTATION")
    add_heading2("4.1 Tools and Technologies")

    # Table 4.1
    table_tools = doc.add_table(rows=7, cols=3)
    table_tools.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_widths = [Inches(1.8), Inches(1.5), Inches(3.2)]
    t_headers = ["Tool / Technology", "Version", "Purpose"]
    for i, h in enumerate(t_headers):
        cell = table_tools.cell(0, i)
        cell.width = t_widths[i]
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(10)

    t_rows = [
        ("Python", "3.12.5", "Programming language"),
        ("Google Colab / VS Code", "Desktop / Cloud", "Development environment"),
        ("Pandas / NumPy", "2.x / 1.x", "Data handling and matrix computations"),
        ("Scikit-learn", "1.x", "Preprocessing and regression models"),
        ("Matplotlib / Seaborn", "3.x / 0.13", "Visualization (heatmaps, distributions, residual plots)"),
        ("Streamlit", "1.40+", "User interface / Multi-page web framework")
    ]
    for r_idx, row in enumerate(t_rows, start=1):
        for c_idx, val in enumerate(row):
            cell = table_tools.cell(r_idx, c_idx)
            cell.width = t_widths[c_idx]
            set_cell_margins(cell, 60, 60, 100, 100)
            if r_idx % 2 == 0:
                set_cell_background(cell, "F8FAFC")
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            if c_idx == 0:
                run.bold = True

    add_heading2("4.2 Module-wise Implementation")
    add_p("● Dataset loading and exploration module: Ingests 1,515 property records and displays statistical profile.")
    add_p("● Data preprocessing module: SimpleImputer (median/mode) and OneHotEncoder pipelines.")
    add_p("● Feature extraction/selection module: Correlation heatmap and variance analysis.")
    add_p("● Model training and evaluation module: GridSearchCV with 5-fold cross-validation across 5 algorithms.")
    add_p("● Best model saving/loading module: joblib serialization (best_model.joblib).")
    add_p("● GUI input and prediction module: Interactive multi-page Streamlit interface with live inference.")

    # Page Break to Page 4
    doc.add_page_break()

    # =========================================================================
    # PAGE 4: SECTIONS 4.3, 5, 6, 7, 8, 9
    # =========================================================================
    add_heading2("4.3 GUI Implementation")
    add_p(
        "The GUI is built using modern Streamlit multi-page architecture. It contains: (1) Market Analytics (interactive boxplots and correlation heatmap); (2) Model Governance (leaderboard table and residual diagnostic plots); and (3) Property Appraisal (dual-column sliders, 5 archetype presets, fair market valuation with +-1 RMSE confidence interval, and real-time What-If CapEx simulator)."
    )

    add_heading1("5. COMPLETE CODE")
    add_p(
        "The complete runnable source code is organized into modular Python files: train_models.py (training pipeline), app.py (dashboard), utils.py (utilities), and pages/ (sub-pages). All code is pushed to the linked GitHub repository and Colab notebook."
    )

    add_heading1("6. RESULTS AND DISCUSSION")
    add_heading2("6.1 Required Output Screenshots")

    # Add images
    if os.path.exists("outputs/r2_comparison.png"):
        p_img1 = doc.add_paragraph()
        p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img1.paragraph_format.space_after = Pt(2)
        p_img1.add_run().add_picture("outputs/r2_comparison.png", width=Inches(4.5))
        add_p("Figure 1: Model Comparison (R² across Train, Test, and CV folds)", font_size=10, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)

    if os.path.exists("outputs/actual_vs_predicted.png"):
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.paragraph_format.space_after = Pt(2)
        p_img2.add_run().add_picture("outputs/actual_vs_predicted.png", width=Inches(5.0))
        add_p("Figure 2: Actual vs Predicted Parity Alignment & Residual Homoscedasticity", font_size=10, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

    add_heading2("6.2 Result Analysis")
    add_p(
        "1. Linear Model Bottleneck: Linear, Ridge, and Lasso regression all plateaued at R² approx 0.8869 with an error rate of 29% MAPE, confirming that additive linear combinations miss multiplicative location-typology multipliers.\n"
        "2. Decision Tree vs Ensemble: A single Decision Tree reached 0.9465 R² but exhibited variance between training (0.9699) and test folds. Random Forest eliminated this variance by averaging 100 decorrelated trees, achieving an extraordinary 0.9604 Test R² and reducing MAPE to 8.95%. Residual plots confirm error variance is homoscedastic with zero mean."
    )

    add_heading1("7. CONCLUSION")
    add_p(
        "This project successfully designed and deployed EstatePulse AI, an automated machine learning property valuation platform. Random Forest Regressor was identified as the optimal architecture, explaining 96.04% of total price variance with an average error of ₹7.83 Lakhs. The interactive Streamlit GUI makes these insights accessible in real time. Future work will integrate GIS satellite mapping and interest rate time-series forecasting."
    )

    add_heading1("8. GITHUB LINK")
    add_p("GitHub Repository URL: https://github.com/Runjhun-Suryawanshi/real-estate-project.git")
    add_p("The repository contains complete code files, dataset, trained model bundle (best_model.joblib), Streamlit GUI application, and outputs.")

    add_heading1("9. REFERENCES")
    add_p("[1] Breiman, L. (2001). Random Forests. Machine Learning, 45(1), 5–32.")
    add_p("[2] Pedregosa, F. et al. (2011). Scikit-learn: Machine Learning in Python. JMLR, 12, 2825–2830.")
    add_p("[3] Streamlit Documentation: Building Multi-Page Web Applications (https://docs.streamlit.io).")
    add_p("[4] Indian Real Estate Housing Market Benchmark Datasets and Kaggle Repository.")

    # Page Break to Page 5
    doc.add_page_break()

    # =========================================================================
    # PAGE 5: FINAL SUBMISSION CHECKLIST
    # =========================================================================
    p_chk = add_p("FINAL SUBMISSION CHECKLIST", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=20, space_after=20)

    add_p("● Final project report in PDF format.")
    add_p("● Google Colab notebook (.ipynb).")
    add_p("● Colab notebook exported as PDF.")
    add_p("● GitHub repository containing Python code, dataset or dataset link, trained model file, GUI code and Colab notebook.")

    p_fmt = add_p(
        "Formatting: The first page should contain college/university name, department, logo, project title and student details. From page 2 onwards, use project title in Times New Roman 16 pt, headings/subheadings in Times New Roman 14 pt, and normal text in Times New Roman 12 pt, justified.",
        font_size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=20, space_after=12
    )
    p_chk_note = add_p(
        "Replace all << placeholders >> with actual project information. Verify that all screenshots and numerical results are from your own implementation.",
        font_size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=40
    )

    doc.save(DOCX_FILENAME)
    print(f"Successfully generated: {DOCX_FILENAME}")

if __name__ == "__main__":
    create_docx()
